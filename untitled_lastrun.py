#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
This experiment was created using PsychoPy3 Experiment Builder (v2026.2.3),
    on 九月 30, 2026, at 10:39
If you publish work using this script the most relevant publication is:

    Peirce J, Gray JR, Simpson S, MacAskill M, Höchenberger R, Sogo H, Kastman E, Lindeløv JK. (2019) 
        PsychoPy2: Experiments in behavior made easy Behav Res 51: 195. 
        https://doi.org/10.3758/s13428-018-01193-y

"""

# --- Import packages ---
from psychopy import locale_setup
from psychopy import prefs
from psychopy import plugins
plugins.activatePlugins()
from psychopy import sound, gui, visual, core, data, event, logging, clock, colors, layout, hardware
from psychopy.tools import environmenttools
from psychopy.constants import (
    NOT_STARTED, STARTED, PLAYING, PAUSED, STOPPED, STOPPING, FINISHED, PRESSED, 
    RELEASED, FOREVER, priority
)

import numpy as np  # whole numpy lib is available, prepend 'np.'
from numpy import (sin, cos, tan, log, log10, pi, average,
                   sqrt, std, deg2rad, rad2deg, linspace, asarray)
from numpy.random import random, randint, normal, shuffle, choice as randchoice
import os  # handy system and path functions
import sys  # to get file system encoding

from psychopy.hardware import keyboard

# --- Setup global variables (available in all functions) ---
# create a device manager to handle hardware (keyboards, mice, mirophones, speakers, etc.)
deviceManager = hardware.DeviceManager()
# ensure that relative paths start from the same directory as this script
_thisDir = os.path.dirname(os.path.abspath(__file__))
# store info about the experiment session
psychopyVersion = '2026.2.3'
expName = 'untitled'  # from the Builder filename that created this script
expVersion = ''
# a list of functions to run when the experiment ends (starts off blank)
runAtExit = []
# information about this experiment
expInfo = {
    'participant': f"{randint(0, 999999):06.0f}",
    'session': '001',
    'date|hid': data.getDateStr(),
    'expName|hid': expName,
    'expVersion|hid': expVersion,
    'psychopyVersion|hid': psychopyVersion,
}

# --- Define some variables which will change depending on pilot mode ---
'''
To run in pilot mode, either use the run/pilot toggle in Builder, Coder and Runner, 
or run the experiment with `--pilot` as an argument. To change what pilot 
#mode does, check out the 'Pilot mode' tab in preferences.
'''
# work out from system args whether we are running in pilot mode
PILOTING = core.setPilotModeFromArgs()
# start off with values from experiment settings
_fullScr = True
_winSize = (1024, 768)
# if in pilot mode, apply overrides according to preferences
if PILOTING:
    # force windowed mode
    if prefs.piloting['forceWindowed']:
        _fullScr = False
        # set window size
        _winSize = prefs.piloting['forcedWindowSize']
    # replace default participant ID
    if prefs.piloting['replaceParticipantID']:
        expInfo['participant'] = 'pilot'

def showExpInfoDlg(expInfo):
    """
    Show participant info dialog.
    Parameters
    ==========
    expInfo : dict
        Information about this experiment.
    
    Returns
    ==========
    dict
        Information about this experiment.
    """
    # show participant info dialog
    dlg = gui.DlgFromDict(
        dictionary=expInfo, sortKeys=False, title=expName, alwaysOnTop=True
    )
    if dlg.OK == False:
        core.quit()  # user pressed cancel
    # return expInfo
    return expInfo


def setupData(expInfo, dataDir=None):
    """
    Make an ExperimentHandler to handle trials and saving.
    
    Parameters
    ==========
    expInfo : dict
        Information about this experiment, created by the `setupExpInfo` function.
    dataDir : Path, str or None
        Folder to save the data to, leave as None to create a folder in the current directory.    
    Returns
    ==========
    psychopy.data.ExperimentHandler
        Handler object for this experiment, contains the data to save and information about 
        where to save it to.
    """
    # remove dialog-specific syntax from expInfo
    for key, val in expInfo.copy().items():
        newKey, _ = data.utils.parsePipeSyntax(key)
        expInfo[newKey] = expInfo.pop(key)
    
    # data file name stem = absolute path + name; later add .psyexp, .csv, .log, etc
    if dataDir is None:
        dataDir = _thisDir
    filename = 'data/' + expInfo['participant'] + '_' + expName + '_' + expInfo['date']
    # make sure filename is relative to dataDir
    if os.path.isabs(filename):
        dataDir = os.path.commonprefix([dataDir, filename])
        filename = os.path.relpath(filename, dataDir)
    
    # an ExperimentHandler isn't essential but helps with data saving
    thisExp = data.ExperimentHandler(
        name=expName, version=expVersion,
        extraInfo=expInfo, runtimeInfo=None,
        originPath='C:\\Users\\27726\\Desktop\\资料\\课程学习\\大三上\\认知科学\\实验2\\untitled_lastrun.py',
        savePickle=True, saveWideText=True,
        dataFileName=dataDir + os.sep + filename, sortColumns='time'
    )
    # store pilot mode in data file
    thisExp.addData('piloting', PILOTING, priority=priority.LOW)
    thisExp.setPriority('thisRow.t', priority.CRITICAL)
    thisExp.setPriority('expName', priority.LOW)
    # return experiment handler
    return thisExp


def setupLogging(filename):
    """
    Setup a log file and tell it what level to log at.
    
    Parameters
    ==========
    filename : str or pathlib.Path
        Filename to save log file and data files as, doesn't need an extension.
    
    Returns
    ==========
    psychopy.logging.LogFile
        Text stream to receive inputs from the logging system.
    """
    # set how much information should be printed to the console / app
    if PILOTING:
        logging.console.setLevel(
            prefs.piloting['pilotConsoleLoggingLevel']
        )
    else:
        logging.console.setLevel('warning')
    # save a log file for detail verbose info
    logFile = logging.LogFile(filename+'.log')
    if PILOTING:
        logFile.setLevel(
            prefs.piloting['pilotLoggingLevel']
        )
    else:
        logFile.setLevel(
            logging.getLevel('info')
        )
    
    return logFile


def setupWindow(expInfo=None, win=None):
    """
    Setup the Window
    
    Parameters
    ==========
    expInfo : dict
        Information about this experiment, created by the `setupExpInfo` function.
    win : psychopy.visual.Window
        Window to setup - leave as None to create a new window.
    
    Returns
    ==========
    psychopy.visual.Window
        Window in which to run this experiment.
    """
    if PILOTING:
        logging.debug('Fullscreen settings ignored as running in pilot mode.')
    
    if win is None:
        # if not given a window to setup, make one
        win = visual.Window(
            size=_winSize, fullscr=_fullScr, screen=0,
            winType='pyglet', allowGUI=False, allowStencil=True,
            monitor='testMonitor', color=[0,0,0], colorSpace='rgb',
            backgroundImage='', backgroundFit='none',
            blendMode='avg', useFBO=True,
            units='height',
            checkTiming=False  # we're going to do this ourselves in a moment
        )
    else:
        # if we have a window, just set the attributes which are safe to set
        win.color = [0,0,0]
        win.colorSpace = 'rgb'
        win.backgroundImage = ''
        win.backgroundFit = 'none'
        win.units = 'height'
    if expInfo is not None:
        # get/measure frame rate if not already in expInfo
        if win._monitorFrameRate is None:
            win._monitorFrameRate = win.getActualFrameRate(infoMsg='Attempting to measure frame rate of screen, please wait...')
        expInfo['frameRate'] = win._monitorFrameRate
    win.hideMessage()
    if PILOTING:
        # show a visual indicator if we're in piloting mode
        if prefs.piloting['showPilotingIndicator']:
            win.showPilotingIndicator()
        # always show the mouse in piloting mode
        if prefs.piloting['forceMouseVisible']:
            win.mouseVisible = True
    
    return win


def setupDevices(expInfo, thisExp, win):
    """
    Setup whatever devices are available (mouse, keyboard, speaker, eyetracker, etc.) and add them to 
    the device manager (deviceManager)
    
    Parameters
    ==========
    expInfo : dict
        Information about this experiment, created by the `setupExpInfo` function.
    thisExp : psychopy.data.ExperimentHandler
        Handler object for this experiment, contains the data to save and information about 
        where to save it to.
    win : psychopy.visual.Window
        Window in which to run this experiment.
    Returns
    ==========
    bool
        True if completed successfully.
    """
    # --- Setup input devices ---
    ioConfig = {}
    ioSession = ioServer = eyetracker = None
    
    # store ioServer object in the device manager
    deviceManager.ioServer = ioServer
    
    # create a default keyboard (e.g. to check for escape)
    if deviceManager.getDevice('defaultKeyboard') is None:
        deviceManager.addDevice(
            deviceClass='keyboard', deviceName='defaultKeyboard', backend='ptb'
        )
    # return True if completed successfully
    return True

def pauseExperiment(thisExp, win=None, timers=[], currentRoutine=None):
    """
    Pause this experiment, preventing the flow from advancing to the next routine until resumed.
    
    Parameters
    ==========
    thisExp : psychopy.data.ExperimentHandler
        Handler object for this experiment, contains the data to save and information about 
        where to save it to.
    win : psychopy.visual.Window
        Window for this experiment.
    timers : list, tuple
        List of timers to reset once pausing is finished.
    currentRoutine : psychopy.data.Routine
        Current Routine we are in at time of pausing, if any. This object tells PsychoPy what Components to pause/play/dispatch.
    """
    # if we are not paused, do nothing
    if thisExp.status != PAUSED:
        return
    
    # start a timer to figure out how long we're paused for
    pauseTimer = core.Clock()
    # pause any playback components
    if currentRoutine is not None:
        for comp in currentRoutine.getPlaybackComponents():
            comp.pause()
    # make sure we have a keyboard
    defaultKeyboard = deviceManager.getDevice('defaultKeyboard')
    if defaultKeyboard is None:
        defaultKeyboard = deviceManager.addKeyboard(
            deviceClass='keyboard',
            deviceName='defaultKeyboard',
            backend='PsychToolbox',
        )
    # run a while loop while we wait to unpause
    while thisExp.status == PAUSED:
        # check for quit (typically the Esc key)
        if defaultKeyboard.getKeys(keyList=['escape']):
            endExperiment(thisExp, win=win)
        # dispatch messages on response components
        if currentRoutine is not None:
            for comp in currentRoutine.getDispatchComponents():
                comp.device.dispatchMessages()
        # sleep 1ms so other threads can execute
        clock.time.sleep(0.001)
    # if stop was requested while paused, quit
    if thisExp.status == FINISHED:
        endExperiment(thisExp, win=win)
    # resume any playback components
    if currentRoutine is not None:
        for comp in currentRoutine.getPlaybackComponents():
            comp.play()
    # reset any timers
    for timer in timers:
        timer.addTime(-pauseTimer.getTime())


def run(expInfo, thisExp, win, globalClock=None, thisSession=None):
    """
    Run the experiment flow.
    
    Parameters
    ==========
    expInfo : dict
        Information about this experiment, created by the `setupExpInfo` function.
    thisExp : psychopy.data.ExperimentHandler
        Handler object for this experiment, contains the data to save and information about 
        where to save it to.
    psychopy.visual.Window
        Window in which to run this experiment.
    globalClock : psychopy.core.clock.Clock or None
        Clock to get global time from - supply None to make a new one.
    thisSession : psychopy.session.Session or None
        Handle of the Session object this experiment is being run from, if any.
    """
    # mark experiment as started
    thisExp.status = STARTED
    # update experiment info
    expInfo['date'] = data.getDateStr()
    expInfo['expName'] = expName
    expInfo['expVersion'] = expVersion
    expInfo['psychopyVersion'] = psychopyVersion
    # make sure window is set to foreground to prevent losing focus
    win.winHandle.activate()
    # make sure variables created by exec are available globally
    exec = environmenttools.setExecEnvironment(globals())
    # get device handles from dict of input devices
    ioServer = deviceManager.ioServer
    # get/create a default keyboard (e.g. to check for escape)
    defaultKeyboard = deviceManager.getDevice('defaultKeyboard')
    if defaultKeyboard is None:
        deviceManager.addDevice(
            deviceClass='keyboard', deviceName='defaultKeyboard', backend='PsychToolbox'
        )
    eyetracker = deviceManager.getDevice('eyetracker')
    # make sure we're running in the directory for this experiment
    os.chdir(_thisDir)
    # get filename from ExperimentHandler for convenience
    filename = thisExp.dataFileName
    frameTolerance = 0.001  # how close to onset before 'same' frame
    endExpNow = False  # flag for 'escape' or other condition => quit the exp
    # get frame duration from frame rate in expInfo
    if 'frameRate' in expInfo and expInfo['frameRate'] is not None:
        frameDur = 1.0 / round(expInfo['frameRate'])
    else:
        frameDur = 1.0 / 60.0  # could not measure, so guess
    
    # Start Code - component code to be run after the window creation
    
    # --- Initialize components for Routine "HELLO" ---
    SETTING_IMAGE_1 = visual.ImageStim(
        win=win,
        name='SETTING_IMAGE_1', units='norm', 
        image='assets/setting.png', mask=None, anchor='center',
        ori=0.0, pos=(0, 0), draggable=False, size=(2, 2),
        color=[1,1,1], colorSpace='rgb', opacity=0.9,
        flipHoriz=False, flipVert=False,
        texRes=128.0, interpolate=True, depth=0.0)
    CONTINUE_KEY_1 = keyboard.Keyboard(deviceName='defaultKeyboard', backend='PsychToolbox')
    HELLO_TEXT = visual.TextStim(win=win, name='HELLO_TEXT',
        text='欢迎参加本次实验！',
        font='Arial',
        pos=(0, 0), draggable=False, height=0.1, wrapWidth=None, ori=0.0, 
        color=(-1.0000, -1.0000, -1.0000), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-2.0);
    DEMOSTRATION_TEXT = visual.TextStim(win=win, name='DEMOSTRATION_TEXT',
        text='（按下 SPACE 键以了解实验规则）\n\n\n温馨提示：切换英文输入法体验更佳',
        font='Arial',
        pos=(0, -0.3), draggable=False, height=0.05, wrapWidth=None, ori=0.0, 
        color=(-0.4902, -0.5059, -0.4902), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-3.0);
    
    # --- Initialize components for Routine "RULES" ---
    SETTING_IMAGE_2 = visual.ImageStim(
        win=win,
        name='SETTING_IMAGE_2', units='norm', 
        image='assets/setting.png', mask=None, anchor='center',
        ori=0.0, pos=(0, 0), draggable=False, size=(2, 2),
        color=[1,1,1], colorSpace='rgb', opacity=0.9,
        flipHoriz=False, flipVert=False,
        texRes=128.0, interpolate=True, depth=0.0)
    CONTINUE_KEY_2 = keyboard.Keyboard(deviceName='defaultKeyboard', backend='PsychToolbox')
    RULE_TEXT_TITLE = visual.TextStim(win=win, name='RULE_TEXT_TITLE',
        text='规则',
        font='Arial',
        pos=(0, 0.3), draggable=False, height=0.1, wrapWidth=None, ori=0.0, 
        color=(-1.0000, -1.0000, -0.9922), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-2.0);
    RULE_TEXT_1 = visual.TextStim(win=win, name='RULE_TEXT_1',
        text='1. 本次实验分为16个循环，每个循环的流程固定',
        font='Arial',
        pos=(-0.4, 0.2), draggable=False, height=0.04, wrapWidth=None, ori=0.0, 
        color=(-1.0000, -1.0000, -1.0000), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-3.0);
    RULE_TEXT_2 = visual.TextStim(win=win, name='RULE_TEXT_2',
        text='2.循环开始，屏幕中心会出现一个红色小球，你需要盯着这个小球',
        font='Arial',
        pos=(-0.4, 0.1), draggable=False, height=0.04, wrapWidth=None, ori=0.0, 
        color=(-1.0000, -1.0000, -1.0000), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-4.0);
    RULE_TEXT_3 = visual.TextStim(win=win, name='RULE_TEXT_3',
        text='3.小球消失的瞬间，你需要按下空格键',
        font='Arial',
        pos=(-0.4, 0.0), draggable=False, height=0.04, wrapWidth=None, ori=0.0, 
        color=(-1.0000, -1.0000, -1.0000), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-5.0);
    RULE_TEXT_4 = visual.TextStim(win=win, name='RULE_TEXT_4',
        text='4.你按下空格键的同时，屏幕上会出现人或物的图片（可能倒立）',
        font='Arial',
        pos=(-0.4, -0.1), draggable=False, height=0.04, wrapWidth=None, ori=0.0, 
        color=(-1.0000, -1.0000, -1.0000), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-6.0);
    RULE_TEXT_5 = visual.TextStim(win=win, name='RULE_TEXT_5',
        text='5.你需要在辨认出人/物的瞬间按下SPACE',
        font='Arial',
        pos=(-0.4, -0.2), draggable=False, height=0.04, wrapWidth=None, ori=0.0, 
        color=(-1.0000, -1.0000, -1.0000), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-7.0);
    RULE_TEXT_6 = visual.TextStim(win=win, name='RULE_TEXT_6',
        text='6.循环最后你需要写出你刚刚看到的人/物的名字（最好英文）\n（不要求完全正确，如ikun\\caixukun\\kunkun都算对）',
        font='Arial',
        pos=(-0.4, -0.3), draggable=False, height=0.04, wrapWidth=None, ori=0.0, 
        color=(-1.0000, -1.0000, -1.0000), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-8.0);
    RULE_TEXT_7 = visual.TextStim(win=win, name='RULE_TEXT_7',
        text='7.你需要坐起来做实验，准备好请按下SPACE',
        font='Arial',
        pos=(-0.4, -0.4), draggable=False, height=0.04, wrapWidth=None, ori=0.0, 
        color=(-1.0000, -1.0000, -1.0000), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-9.0);
    
    # --- Initialize components for Routine "POINT" ---
    SETTIGN_IMAGE_3 = visual.ImageStim(
        win=win,
        name='SETTIGN_IMAGE_3', units='norm', 
        image='assets/setting.png', mask=None, anchor='center',
        ori=0.0, pos=(0, 0), draggable=False, size=(2, 2),
        color=[1,1,1], colorSpace='rgb', opacity=0.9,
        flipHoriz=False, flipVert=False,
        texRes=128.0, interpolate=True, depth=0.0)
    CONTINUE_KEY_3 = keyboard.Keyboard(deviceName='defaultKeyboard', backend='PsychToolbox')
    Point = visual.ShapeStim(
        win=win, name='Point',
        size=(0.1, 0.1), vertices='circle',
        ori=0.0, pos=(0, 0), draggable=False, anchor='center',
        lineWidth=1.0,
        colorSpace='rgb', lineColor='white', fillColor=(1.0000, -1.0000, -1.0000),
        opacity=None, depth=-2.0, interpolate=True)
    
    # --- Initialize components for Routine "STIMULATE" ---
    SETTING_IMAGE_4 = visual.ImageStim(
        win=win,
        name='SETTING_IMAGE_4', units='norm', 
        image='assets/setting.png', mask=None, anchor='center',
        ori=0.0, pos=(0, 0), draggable=False, size=(2, 2),
        color=[1,1,1], colorSpace='rgb', opacity=0.9,
        flipHoriz=False, flipVert=False,
        texRes=128.0, interpolate=True, depth=0.0)
    CONTINUE_KEY_4 = keyboard.Keyboard(deviceName='defaultKeyboard', backend='PsychToolbox')
    STIMULATE_IMAGE = visual.ImageStim(
        win=win,
        name='STIMULATE_IMAGE', 
        image='default.png', mask=None, anchor='center',
        ori=0.0, pos=(0, 0), draggable=False, size=(0.5, 0.5),
        color=[1,1,1], colorSpace='rgb', opacity=None,
        flipHoriz=False, flipVert=False,
        texRes=128.0, interpolate=True, depth=-2.0)
    
    # --- Initialize components for Routine "RECORD" ---
    SETTING_IMAGE_5 = visual.ImageStim(
        win=win,
        name='SETTING_IMAGE_5', units='norm', 
        image='assets/setting.png', mask=None, anchor='center',
        ori=0.0, pos=(0, 0), draggable=False, size=(2, 2),
        color=[1,1,1], colorSpace='rgb', opacity=0.9,
        flipHoriz=False, flipVert=False,
        texRes=128.0, interpolate=True, depth=0.0)
    TITLE_TEXT = visual.TextStim(win=win, name='TITLE_TEXT',
        text='你刚刚看到了什么？',
        font='Arial',
        pos=(0, 0.3), draggable=False, height=0.1, wrapWidth=None, ori=0.0, 
        color=(-1.0000, -1.0000, -1.0000), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-1.0);
    INPUT = visual.TextBox2(
         win, text=None, placeholder='Type here...', font='Arial',
         ori=0.0, pos=(0, 0), draggable=False,      letterHeight=0.05,
         size=(0.5, 0.5), borderWidth=2.0,
         color=(-1.0000, -1.0000, -1.0000), colorSpace='rgb',
         opacity=None,
         bold=False, italic=False,
         lineSpacing=1.0, speechPoint=None,
         padding=0.0, alignment='center',
         anchor='center', overflow='visible',
         fillColor=None, borderColor=None,
         flipHoriz=False, flipVert=False, 
         languageStyle='LTR',
         formattingSyntax='md',
         editable=True,
         name='INPUT',
         depth=-2, autoLog=True,
    )
    TEXT = visual.TextStim(win=win, name='TEXT',
        text='完成',
        font='Arial',
        pos=(0, -0.3), draggable=False, height=0.05, wrapWidth=None, ori=0.0, 
        color='white', colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-3.0);
    SUBMMIT_KEY = keyboard.Keyboard(deviceName='defaultKeyboard', backend='PsychToolbox')
    
    # --- Initialize components for Routine "THANKS" ---
    THANKS_TEXT = visual.TextStim(win=win, name='THANKS_TEXT',
        text='感恩！',
        font='Arial',
        pos=(0, 0), draggable=False, height=0.1, wrapWidth=None, ori=0.0, 
        color=(-1.0000, -1.0000, -1.0000), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=0.0);
    THANKS_DETAILED_TEXT = visual.TextStim(win=win, name='THANKS_DETAILED_TEXT',
        text='王梓能够完成作业离不开你的支持\n\n（按下空格以结束）',
        font='Arial',
        pos=(0, -0.3), draggable=False, height=0.05, wrapWidth=None, ori=0.0, 
        color=(-0.4902, -0.5059, -0.4902), colorSpace='rgb', opacity=None, 
        languageStyle='LTR',
        depth=-1.0);
    CONTINUE_KEY_5 = keyboard.Keyboard(deviceName='defaultKeyboard', backend='PsychToolbox')
    
    # create some handy timers
    
    # global clock to track the time since experiment started
    if globalClock is None:
        # create a clock if not given one
        globalClock = core.Clock()
    if isinstance(globalClock, str):
        # if given a string, make a clock accoridng to it
        if globalClock == 'float':
            # get timestamps as a simple value
            globalClock = core.Clock(format='float')
        elif globalClock == 'iso':
            # get timestamps in ISO format
            globalClock = core.Clock(format='%Y-%m-%d_%H:%M:%S.%f%z')
        else:
            # get timestamps in a custom format
            globalClock = core.Clock(format=globalClock)
    if ioServer is not None:
        ioServer.syncClock(globalClock)
    logging.setDefaultClock(globalClock)
    if eyetracker is not None:
        eyetracker.enableEventReporting()
    # routine timer to track time remaining of each (possibly non-slip) routine
    routineTimer = core.Clock()
    win.flip()  # flip window to reset last flip timer
    # store the exact time the global clock started
    expInfo['expStart'] = data.getDateStr(
        format='%Y-%m-%d %Hh%M.%S.%f %z', fractionalSecondDigits=6
    )
    
    # --- Prepare to start Routine "HELLO" ---
    # create an object to store info about Routine HELLO
    HELLO = data.Routine(
        name='HELLO',
        components=[SETTING_IMAGE_1, CONTINUE_KEY_1, HELLO_TEXT, DEMOSTRATION_TEXT],
    )
    HELLO.status = NOT_STARTED
    continueRoutine = True
    # update component parameters for each repeat
    # create starting attributes for CONTINUE_KEY_1
    CONTINUE_KEY_1.keys = []
    CONTINUE_KEY_1.rt = []
    _CONTINUE_KEY_1_allKeys = []
    # store start times for HELLO
    HELLO.tStartRefresh = win.getFutureFlipTime(clock=globalClock)
    HELLO.tStart = globalClock.getTime(format='float')
    HELLO.status = STARTED
    thisExp.addData('HELLO.started', HELLO.tStart)
    HELLO.maxDuration = None
    # keep track of which components have finished
    HELLOComponents = HELLO.components
    for thisComponent in HELLO.components:
        thisComponent.tStart = None
        thisComponent.tStop = None
        thisComponent.tStartRefresh = None
        thisComponent.tStopRefresh = None
        if hasattr(thisComponent, 'status'):
            thisComponent.status = NOT_STARTED
    # reset timers
    t = 0
    _timeToFirstFrame = win.getFutureFlipTime(clock="now")
    frameN = -1
    
    # --- Run Routine "HELLO" ---
    thisExp.currentRoutine = HELLO
    HELLO.forceEnded = routineForceEnded = not continueRoutine
    while continueRoutine:
        # get current time
        t = routineTimer.getTime()
        tThisFlip = win.getFutureFlipTime(clock=routineTimer)
        tThisFlipGlobal = win.getFutureFlipTime(clock=None)
        frameN = frameN + 1  # number of completed frames (so 0 is the first frame)
        # update/draw components on each frame
        
        # *SETTING_IMAGE_1* updates
        
        # if SETTING_IMAGE_1 is starting this frame...
        if SETTING_IMAGE_1.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            SETTING_IMAGE_1.frameNStart = frameN  # exact frame index
            SETTING_IMAGE_1.tStart = t  # local t and not account for scr refresh
            SETTING_IMAGE_1.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(SETTING_IMAGE_1, 'tStartRefresh')  # time at next scr refresh
            # update status
            SETTING_IMAGE_1.status = STARTED
            SETTING_IMAGE_1.setAutoDraw(True)
        
        # if SETTING_IMAGE_1 is active this frame...
        if SETTING_IMAGE_1.status == STARTED:
            # update params
            pass
        
        # *CONTINUE_KEY_1* updates
        
        # if CONTINUE_KEY_1 is starting this frame...
        if CONTINUE_KEY_1.status == NOT_STARTED and t >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            CONTINUE_KEY_1.frameNStart = frameN  # exact frame index
            CONTINUE_KEY_1.tStart = t  # local t and not account for scr refresh
            CONTINUE_KEY_1.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(CONTINUE_KEY_1, 'tStartRefresh')  # time at next scr refresh
            # update status
            CONTINUE_KEY_1.status = STARTED
            # keyboard checking is just starting
            CONTINUE_KEY_1.clock.reset()  # now t=0
        if CONTINUE_KEY_1.status == STARTED:
            theseKeys = CONTINUE_KEY_1.getKeys(keyList=['space'], ignoreKeys=["escape"], waitRelease=False)
            _CONTINUE_KEY_1_allKeys.extend(theseKeys)
            if len(_CONTINUE_KEY_1_allKeys):
                CONTINUE_KEY_1.keys = _CONTINUE_KEY_1_allKeys[-1].name  # just the last key pressed
                CONTINUE_KEY_1.rt = _CONTINUE_KEY_1_allKeys[-1].rt
                CONTINUE_KEY_1.duration = _CONTINUE_KEY_1_allKeys[-1].duration
                # a response ends the routine
                continueRoutine = False
        
        # *HELLO_TEXT* updates
        
        # if HELLO_TEXT is starting this frame...
        if HELLO_TEXT.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            HELLO_TEXT.frameNStart = frameN  # exact frame index
            HELLO_TEXT.tStart = t  # local t and not account for scr refresh
            HELLO_TEXT.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(HELLO_TEXT, 'tStartRefresh')  # time at next scr refresh
            # update status
            HELLO_TEXT.status = STARTED
            HELLO_TEXT.setAutoDraw(True)
        
        # if HELLO_TEXT is active this frame...
        if HELLO_TEXT.status == STARTED:
            # update params
            pass
        
        # *DEMOSTRATION_TEXT* updates
        
        # if DEMOSTRATION_TEXT is starting this frame...
        if DEMOSTRATION_TEXT.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            DEMOSTRATION_TEXT.frameNStart = frameN  # exact frame index
            DEMOSTRATION_TEXT.tStart = t  # local t and not account for scr refresh
            DEMOSTRATION_TEXT.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(DEMOSTRATION_TEXT, 'tStartRefresh')  # time at next scr refresh
            # update status
            DEMOSTRATION_TEXT.status = STARTED
            DEMOSTRATION_TEXT.setAutoDraw(True)
        
        # if DEMOSTRATION_TEXT is active this frame...
        if DEMOSTRATION_TEXT.status == STARTED:
            # update params
            pass
        
        # check for quit (typically the Esc key)
        if defaultKeyboard.getKeys(keyList=["escape"]):
            thisExp.status = FINISHED
        if thisExp.status == FINISHED or endExpNow:
            endExperiment(thisExp, win=win)
            return
        # pause experiment here if requested
        if thisExp.status == PAUSED:
            pauseExperiment(
                thisExp=thisExp, 
                win=win, 
                timers=[routineTimer, globalClock], 
                currentRoutine=HELLO,
            )
            # skip the frame we paused on
            continue
        
        # has a Component requested the Routine to end?
        if not continueRoutine:
            HELLO.forceEnded = routineForceEnded = True
        # has the Routine been forcibly ended?
        if HELLO.forceEnded or routineForceEnded:
            break
        # has every Component finished?
        continueRoutine = False
        for thisComponent in HELLO.components:
            if hasattr(thisComponent, "status") and thisComponent.status != FINISHED:
                continueRoutine = True
                break  # at least one component has not yet finished
        
        # refresh the screen
        if continueRoutine:  # don't flip if this routine is over or we'll get a blank screen
            win.flip()
    
    # --- Ending Routine "HELLO" ---
    for thisComponent in HELLO.components:
        if hasattr(thisComponent, "setAutoDraw"):
            thisComponent.setAutoDraw(False)
    # store stop times for HELLO
    HELLO.tStop = globalClock.getTime(format='float')
    HELLO.tStopRefresh = tThisFlipGlobal
    thisExp.addData('HELLO.stopped', HELLO.tStop)
    thisExp.nextEntry()
    # the Routine "HELLO" was not non-slip safe, so reset the non-slip timer
    routineTimer.reset()
    
    # --- Prepare to start Routine "RULES" ---
    # create an object to store info about Routine RULES
    RULES = data.Routine(
        name='RULES',
        components=[SETTING_IMAGE_2, CONTINUE_KEY_2, RULE_TEXT_TITLE, RULE_TEXT_1, RULE_TEXT_2, RULE_TEXT_3, RULE_TEXT_4, RULE_TEXT_5, RULE_TEXT_6, RULE_TEXT_7],
    )
    RULES.status = NOT_STARTED
    continueRoutine = True
    # update component parameters for each repeat
    # create starting attributes for CONTINUE_KEY_2
    CONTINUE_KEY_2.keys = []
    CONTINUE_KEY_2.rt = []
    _CONTINUE_KEY_2_allKeys = []
    # Run 'Begin Routine' code from code
    rule_texts = [
        RULE_TEXT_1,
        RULE_TEXT_2,
        RULE_TEXT_3,
        RULE_TEXT_4,
        RULE_TEXT_5,
        RULE_TEXT_6,
        RULE_TEXT_7
    ]
    
    for t in rule_texts:
        t.anchorHoriz = 'left'
        t.alignText = 'left'
        t.pos = (-0.64, t.pos[1])
    # store start times for RULES
    RULES.tStartRefresh = win.getFutureFlipTime(clock=globalClock)
    RULES.tStart = globalClock.getTime(format='float')
    RULES.status = STARTED
    thisExp.addData('RULES.started', RULES.tStart)
    RULES.maxDuration = None
    # keep track of which components have finished
    RULESComponents = RULES.components
    for thisComponent in RULES.components:
        thisComponent.tStart = None
        thisComponent.tStop = None
        thisComponent.tStartRefresh = None
        thisComponent.tStopRefresh = None
        if hasattr(thisComponent, 'status'):
            thisComponent.status = NOT_STARTED
    # reset timers
    t = 0
    _timeToFirstFrame = win.getFutureFlipTime(clock="now")
    frameN = -1
    
    # --- Run Routine "RULES" ---
    thisExp.currentRoutine = RULES
    RULES.forceEnded = routineForceEnded = not continueRoutine
    while continueRoutine:
        # get current time
        t = routineTimer.getTime()
        tThisFlip = win.getFutureFlipTime(clock=routineTimer)
        tThisFlipGlobal = win.getFutureFlipTime(clock=None)
        frameN = frameN + 1  # number of completed frames (so 0 is the first frame)
        # update/draw components on each frame
        
        # *SETTING_IMAGE_2* updates
        
        # if SETTING_IMAGE_2 is starting this frame...
        if SETTING_IMAGE_2.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            SETTING_IMAGE_2.frameNStart = frameN  # exact frame index
            SETTING_IMAGE_2.tStart = t  # local t and not account for scr refresh
            SETTING_IMAGE_2.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(SETTING_IMAGE_2, 'tStartRefresh')  # time at next scr refresh
            # update status
            SETTING_IMAGE_2.status = STARTED
            SETTING_IMAGE_2.setAutoDraw(True)
        
        # if SETTING_IMAGE_2 is active this frame...
        if SETTING_IMAGE_2.status == STARTED:
            # update params
            pass
        
        # *CONTINUE_KEY_2* updates
        
        # if CONTINUE_KEY_2 is starting this frame...
        if CONTINUE_KEY_2.status == NOT_STARTED and t >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            CONTINUE_KEY_2.frameNStart = frameN  # exact frame index
            CONTINUE_KEY_2.tStart = t  # local t and not account for scr refresh
            CONTINUE_KEY_2.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(CONTINUE_KEY_2, 'tStartRefresh')  # time at next scr refresh
            # update status
            CONTINUE_KEY_2.status = STARTED
            # keyboard checking is just starting
            CONTINUE_KEY_2.clock.reset()  # now t=0
        if CONTINUE_KEY_2.status == STARTED:
            theseKeys = CONTINUE_KEY_2.getKeys(keyList=['space'], ignoreKeys=["escape"], waitRelease=False)
            _CONTINUE_KEY_2_allKeys.extend(theseKeys)
            if len(_CONTINUE_KEY_2_allKeys):
                CONTINUE_KEY_2.keys = _CONTINUE_KEY_2_allKeys[-1].name  # just the last key pressed
                CONTINUE_KEY_2.rt = _CONTINUE_KEY_2_allKeys[-1].rt
                CONTINUE_KEY_2.duration = _CONTINUE_KEY_2_allKeys[-1].duration
                # a response ends the routine
                continueRoutine = False
        
        # *RULE_TEXT_TITLE* updates
        
        # if RULE_TEXT_TITLE is starting this frame...
        if RULE_TEXT_TITLE.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            RULE_TEXT_TITLE.frameNStart = frameN  # exact frame index
            RULE_TEXT_TITLE.tStart = t  # local t and not account for scr refresh
            RULE_TEXT_TITLE.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(RULE_TEXT_TITLE, 'tStartRefresh')  # time at next scr refresh
            # update status
            RULE_TEXT_TITLE.status = STARTED
            RULE_TEXT_TITLE.setAutoDraw(True)
        
        # if RULE_TEXT_TITLE is active this frame...
        if RULE_TEXT_TITLE.status == STARTED:
            # update params
            pass
        
        # *RULE_TEXT_1* updates
        
        # if RULE_TEXT_1 is starting this frame...
        if RULE_TEXT_1.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            RULE_TEXT_1.frameNStart = frameN  # exact frame index
            RULE_TEXT_1.tStart = t  # local t and not account for scr refresh
            RULE_TEXT_1.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(RULE_TEXT_1, 'tStartRefresh')  # time at next scr refresh
            # update status
            RULE_TEXT_1.status = STARTED
            RULE_TEXT_1.setAutoDraw(True)
        
        # if RULE_TEXT_1 is active this frame...
        if RULE_TEXT_1.status == STARTED:
            # update params
            pass
        
        # *RULE_TEXT_2* updates
        
        # if RULE_TEXT_2 is starting this frame...
        if RULE_TEXT_2.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            RULE_TEXT_2.frameNStart = frameN  # exact frame index
            RULE_TEXT_2.tStart = t  # local t and not account for scr refresh
            RULE_TEXT_2.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(RULE_TEXT_2, 'tStartRefresh')  # time at next scr refresh
            # update status
            RULE_TEXT_2.status = STARTED
            RULE_TEXT_2.setAutoDraw(True)
        
        # if RULE_TEXT_2 is active this frame...
        if RULE_TEXT_2.status == STARTED:
            # update params
            pass
        
        # *RULE_TEXT_3* updates
        
        # if RULE_TEXT_3 is starting this frame...
        if RULE_TEXT_3.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            RULE_TEXT_3.frameNStart = frameN  # exact frame index
            RULE_TEXT_3.tStart = t  # local t and not account for scr refresh
            RULE_TEXT_3.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(RULE_TEXT_3, 'tStartRefresh')  # time at next scr refresh
            # update status
            RULE_TEXT_3.status = STARTED
            RULE_TEXT_3.setAutoDraw(True)
        
        # if RULE_TEXT_3 is active this frame...
        if RULE_TEXT_3.status == STARTED:
            # update params
            pass
        
        # *RULE_TEXT_4* updates
        
        # if RULE_TEXT_4 is starting this frame...
        if RULE_TEXT_4.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            RULE_TEXT_4.frameNStart = frameN  # exact frame index
            RULE_TEXT_4.tStart = t  # local t and not account for scr refresh
            RULE_TEXT_4.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(RULE_TEXT_4, 'tStartRefresh')  # time at next scr refresh
            # update status
            RULE_TEXT_4.status = STARTED
            RULE_TEXT_4.setAutoDraw(True)
        
        # if RULE_TEXT_4 is active this frame...
        if RULE_TEXT_4.status == STARTED:
            # update params
            pass
        
        # *RULE_TEXT_5* updates
        
        # if RULE_TEXT_5 is starting this frame...
        if RULE_TEXT_5.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            RULE_TEXT_5.frameNStart = frameN  # exact frame index
            RULE_TEXT_5.tStart = t  # local t and not account for scr refresh
            RULE_TEXT_5.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(RULE_TEXT_5, 'tStartRefresh')  # time at next scr refresh
            # update status
            RULE_TEXT_5.status = STARTED
            RULE_TEXT_5.setAutoDraw(True)
        
        # if RULE_TEXT_5 is active this frame...
        if RULE_TEXT_5.status == STARTED:
            # update params
            pass
        
        # *RULE_TEXT_6* updates
        
        # if RULE_TEXT_6 is starting this frame...
        if RULE_TEXT_6.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            RULE_TEXT_6.frameNStart = frameN  # exact frame index
            RULE_TEXT_6.tStart = t  # local t and not account for scr refresh
            RULE_TEXT_6.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(RULE_TEXT_6, 'tStartRefresh')  # time at next scr refresh
            # update status
            RULE_TEXT_6.status = STARTED
            RULE_TEXT_6.setAutoDraw(True)
        
        # if RULE_TEXT_6 is active this frame...
        if RULE_TEXT_6.status == STARTED:
            # update params
            pass
        
        # *RULE_TEXT_7* updates
        
        # if RULE_TEXT_7 is starting this frame...
        if RULE_TEXT_7.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            RULE_TEXT_7.frameNStart = frameN  # exact frame index
            RULE_TEXT_7.tStart = t  # local t and not account for scr refresh
            RULE_TEXT_7.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(RULE_TEXT_7, 'tStartRefresh')  # time at next scr refresh
            # update status
            RULE_TEXT_7.status = STARTED
            RULE_TEXT_7.setAutoDraw(True)
        
        # if RULE_TEXT_7 is active this frame...
        if RULE_TEXT_7.status == STARTED:
            # update params
            pass
        
        # check for quit (typically the Esc key)
        if defaultKeyboard.getKeys(keyList=["escape"]):
            thisExp.status = FINISHED
        if thisExp.status == FINISHED or endExpNow:
            endExperiment(thisExp, win=win)
            return
        # pause experiment here if requested
        if thisExp.status == PAUSED:
            pauseExperiment(
                thisExp=thisExp, 
                win=win, 
                timers=[routineTimer, globalClock], 
                currentRoutine=RULES,
            )
            # skip the frame we paused on
            continue
        
        # has a Component requested the Routine to end?
        if not continueRoutine:
            RULES.forceEnded = routineForceEnded = True
        # has the Routine been forcibly ended?
        if RULES.forceEnded or routineForceEnded:
            break
        # has every Component finished?
        continueRoutine = False
        for thisComponent in RULES.components:
            if hasattr(thisComponent, "status") and thisComponent.status != FINISHED:
                continueRoutine = True
                break  # at least one component has not yet finished
        
        # refresh the screen
        if continueRoutine:  # don't flip if this routine is over or we'll get a blank screen
            win.flip()
    
    # --- Ending Routine "RULES" ---
    for thisComponent in RULES.components:
        if hasattr(thisComponent, "setAutoDraw"):
            thisComponent.setAutoDraw(False)
    # store stop times for RULES
    RULES.tStop = globalClock.getTime(format='float')
    RULES.tStopRefresh = tThisFlipGlobal
    thisExp.addData('RULES.stopped', RULES.tStop)
    thisExp.nextEntry()
    # the Routine "RULES" was not non-slip safe, so reset the non-slip timer
    routineTimer.reset()
    
    # set up handler to look after randomisation of conditions etc
    trials = data.TrialHandler2(
        name='trials',
        nReps=1, 
        method='random', 
        extraInfo=expInfo, 
        originPath=-1, 
        trialList=data.importConditions('conditions.xlsx'), 
        seed=None, 
        isTrials=True, 
    )
    thisExp.addLoop(trials)  # add the loop to the experiment
    thisTrial = trials.trialList[0]  # so we can initialise stimuli with some values
    # abbreviate parameter names if possible (e.g. rgb = thisTrial.rgb)
    if thisTrial != None:
        for paramName in thisTrial:
            globals()[paramName] = thisTrial[paramName]
    if thisSession is not None:
        # if running in a Session with a Liaison client, send data up to now
        thisSession.sendExperimentData()
    
    for thisTrial in trials:
        trials.status = STARTED
        if hasattr(thisTrial, 'status'):
            thisTrial.status = STARTED
        currentLoop = trials
        thisExp.timestampOnFlip(win, 'thisRow.t', format=globalClock.format)
        if thisSession is not None:
            # if running in a Session with a Liaison client, send data up to now
            thisSession.sendExperimentData()
        # abbreviate parameter names if possible (e.g. rgb = thisTrial.rgb)
        if thisTrial != None:
            for paramName in thisTrial:
                globals()[paramName] = thisTrial[paramName]
        
        # --- Prepare to start Routine "POINT" ---
        # create an object to store info about Routine POINT
        POINT = data.Routine(
            name='POINT',
            components=[SETTIGN_IMAGE_3, CONTINUE_KEY_3, Point],
        )
        POINT.status = NOT_STARTED
        continueRoutine = True
        # update component parameters for each repeat
        # create starting attributes for CONTINUE_KEY_3
        CONTINUE_KEY_3.keys = []
        CONTINUE_KEY_3.rt = []
        _CONTINUE_KEY_3_allKeys = []
        # Run 'Begin Routine' code from code_2
        point_duration = np.random.uniform(0.7, 1.3)
        # store start times for POINT
        POINT.tStartRefresh = win.getFutureFlipTime(clock=globalClock)
        POINT.tStart = globalClock.getTime(format='float')
        POINT.status = STARTED
        thisExp.addData('POINT.started', POINT.tStart)
        POINT.maxDuration = None
        # keep track of which components have finished
        POINTComponents = POINT.components
        for thisComponent in POINT.components:
            thisComponent.tStart = None
            thisComponent.tStop = None
            thisComponent.tStartRefresh = None
            thisComponent.tStopRefresh = None
            if hasattr(thisComponent, 'status'):
                thisComponent.status = NOT_STARTED
        # reset timers
        t = 0
        _timeToFirstFrame = win.getFutureFlipTime(clock="now")
        frameN = -1
        
        # --- Run Routine "POINT" ---
        thisExp.currentRoutine = POINT
        POINT.forceEnded = routineForceEnded = not continueRoutine
        while continueRoutine:
            # if trial has changed, end Routine now
            if hasattr(thisTrial, 'status') and thisTrial.status == STOPPING:
                continueRoutine = False
            # get current time
            t = routineTimer.getTime()
            tThisFlip = win.getFutureFlipTime(clock=routineTimer)
            tThisFlipGlobal = win.getFutureFlipTime(clock=None)
            frameN = frameN + 1  # number of completed frames (so 0 is the first frame)
            # update/draw components on each frame
            
            # *SETTIGN_IMAGE_3* updates
            
            # if SETTIGN_IMAGE_3 is starting this frame...
            if SETTIGN_IMAGE_3.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
                # keep track of start time/frame for later
                SETTIGN_IMAGE_3.frameNStart = frameN  # exact frame index
                SETTIGN_IMAGE_3.tStart = t  # local t and not account for scr refresh
                SETTIGN_IMAGE_3.tStartRefresh = tThisFlipGlobal  # on global time
                win.timeOnFlip(SETTIGN_IMAGE_3, 'tStartRefresh')  # time at next scr refresh
                # update status
                SETTIGN_IMAGE_3.status = STARTED
                SETTIGN_IMAGE_3.setAutoDraw(True)
            
            # if SETTIGN_IMAGE_3 is active this frame...
            if SETTIGN_IMAGE_3.status == STARTED:
                # update params
                pass
            
            # *CONTINUE_KEY_3* updates
            
            # if CONTINUE_KEY_3 is starting this frame...
            if CONTINUE_KEY_3.status == NOT_STARTED and t >= 0.0-frameTolerance:
                # keep track of start time/frame for later
                CONTINUE_KEY_3.frameNStart = frameN  # exact frame index
                CONTINUE_KEY_3.tStart = t  # local t and not account for scr refresh
                CONTINUE_KEY_3.tStartRefresh = tThisFlipGlobal  # on global time
                win.timeOnFlip(CONTINUE_KEY_3, 'tStartRefresh')  # time at next scr refresh
                # update status
                CONTINUE_KEY_3.status = STARTED
                # keyboard checking is just starting
                CONTINUE_KEY_3.clock.reset()  # now t=0
            if CONTINUE_KEY_3.status == STARTED:
                theseKeys = CONTINUE_KEY_3.getKeys(keyList=['space'], ignoreKeys=["escape"], waitRelease=False)
                _CONTINUE_KEY_3_allKeys.extend(theseKeys)
                if len(_CONTINUE_KEY_3_allKeys):
                    CONTINUE_KEY_3.keys = _CONTINUE_KEY_3_allKeys[-1].name  # just the last key pressed
                    CONTINUE_KEY_3.rt = _CONTINUE_KEY_3_allKeys[-1].rt
                    CONTINUE_KEY_3.duration = _CONTINUE_KEY_3_allKeys[-1].duration
                    # a response ends the routine
                    continueRoutine = False
            
            # *Point* updates
            
            # if Point is starting this frame...
            if Point.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
                # keep track of start time/frame for later
                Point.frameNStart = frameN  # exact frame index
                Point.tStart = t  # local t and not account for scr refresh
                Point.tStartRefresh = tThisFlipGlobal  # on global time
                win.timeOnFlip(Point, 'tStartRefresh')  # time at next scr refresh
                # add timestamp to datafile
                thisExp.timestampOnFlip(win, 'Point.started')
                # update status
                Point.status = STARTED
                Point.setAutoDraw(True)
            
            # if Point is active this frame...
            if Point.status == STARTED:
                # update params
                pass
            
            # if Point is stopping this frame...
            if Point.status == STARTED:
                # is it time to stop? (based on global clock, using actual start)
                if tThisFlipGlobal > Point.tStartRefresh + point_duration-frameTolerance:
                    # keep track of stop time/frame for later
                    Point.tStop = t  # not accounting for scr refresh
                    Point.tStopRefresh = tThisFlipGlobal  # on global time
                    Point.frameNStop = frameN  # exact frame index
                    # add timestamp to datafile
                    thisExp.timestampOnFlip(win, 'Point.stopped')
                    # update status
                    Point.status = FINISHED
                    Point.setAutoDraw(False)
            
            # check for quit (typically the Esc key)
            if defaultKeyboard.getKeys(keyList=["escape"]):
                thisExp.status = FINISHED
            if thisExp.status == FINISHED or endExpNow:
                endExperiment(thisExp, win=win)
                return
            # pause experiment here if requested
            if thisExp.status == PAUSED:
                pauseExperiment(
                    thisExp=thisExp, 
                    win=win, 
                    timers=[routineTimer, globalClock], 
                    currentRoutine=POINT,
                )
                # skip the frame we paused on
                continue
            
            # has a Component requested the Routine to end?
            if not continueRoutine:
                POINT.forceEnded = routineForceEnded = True
            # has the Routine been forcibly ended?
            if POINT.forceEnded or routineForceEnded:
                break
            # has every Component finished?
            continueRoutine = False
            for thisComponent in POINT.components:
                if hasattr(thisComponent, "status") and thisComponent.status != FINISHED:
                    continueRoutine = True
                    break  # at least one component has not yet finished
            
            # refresh the screen
            if continueRoutine:  # don't flip if this routine is over or we'll get a blank screen
                win.flip()
        
        # --- Ending Routine "POINT" ---
        for thisComponent in POINT.components:
            if hasattr(thisComponent, "setAutoDraw"):
                thisComponent.setAutoDraw(False)
        # store stop times for POINT
        POINT.tStop = globalClock.getTime(format='float')
        POINT.tStopRefresh = tThisFlipGlobal
        thisExp.addData('POINT.stopped', POINT.tStop)
        # the Routine "POINT" was not non-slip safe, so reset the non-slip timer
        routineTimer.reset()
        
        # --- Prepare to start Routine "STIMULATE" ---
        # create an object to store info about Routine STIMULATE
        STIMULATE = data.Routine(
            name='STIMULATE',
            components=[SETTING_IMAGE_4, CONTINUE_KEY_4, STIMULATE_IMAGE],
        )
        STIMULATE.status = NOT_STARTED
        continueRoutine = True
        # update component parameters for each repeat
        # create starting attributes for CONTINUE_KEY_4
        CONTINUE_KEY_4.keys = []
        CONTINUE_KEY_4.rt = []
        _CONTINUE_KEY_4_allKeys = []
        STIMULATE_IMAGE.setImage(IMAGES)
        # Run 'Begin Routine' code from code_3
        if UP_SIDE_DOWN:
            STIMULATE_IMAGE.ori = 180
        else:
            STIMULATE_IMAGE.ori = 0
        # store start times for STIMULATE
        STIMULATE.tStartRefresh = win.getFutureFlipTime(clock=globalClock)
        STIMULATE.tStart = globalClock.getTime(format='float')
        STIMULATE.status = STARTED
        thisExp.addData('STIMULATE.started', STIMULATE.tStart)
        STIMULATE.maxDuration = None
        # keep track of which components have finished
        STIMULATEComponents = STIMULATE.components
        for thisComponent in STIMULATE.components:
            thisComponent.tStart = None
            thisComponent.tStop = None
            thisComponent.tStartRefresh = None
            thisComponent.tStopRefresh = None
            if hasattr(thisComponent, 'status'):
                thisComponent.status = NOT_STARTED
        # reset timers
        t = 0
        _timeToFirstFrame = win.getFutureFlipTime(clock="now")
        frameN = -1
        
        # --- Run Routine "STIMULATE" ---
        thisExp.currentRoutine = STIMULATE
        STIMULATE.forceEnded = routineForceEnded = not continueRoutine
        while continueRoutine:
            # if trial has changed, end Routine now
            if hasattr(thisTrial, 'status') and thisTrial.status == STOPPING:
                continueRoutine = False
            # get current time
            t = routineTimer.getTime()
            tThisFlip = win.getFutureFlipTime(clock=routineTimer)
            tThisFlipGlobal = win.getFutureFlipTime(clock=None)
            frameN = frameN + 1  # number of completed frames (so 0 is the first frame)
            # update/draw components on each frame
            
            # *SETTING_IMAGE_4* updates
            
            # if SETTING_IMAGE_4 is starting this frame...
            if SETTING_IMAGE_4.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
                # keep track of start time/frame for later
                SETTING_IMAGE_4.frameNStart = frameN  # exact frame index
                SETTING_IMAGE_4.tStart = t  # local t and not account for scr refresh
                SETTING_IMAGE_4.tStartRefresh = tThisFlipGlobal  # on global time
                win.timeOnFlip(SETTING_IMAGE_4, 'tStartRefresh')  # time at next scr refresh
                # update status
                SETTING_IMAGE_4.status = STARTED
                SETTING_IMAGE_4.setAutoDraw(True)
            
            # if SETTING_IMAGE_4 is active this frame...
            if SETTING_IMAGE_4.status == STARTED:
                # update params
                pass
            
            # *CONTINUE_KEY_4* updates
            waitOnFlip = False
            
            # if CONTINUE_KEY_4 is starting this frame...
            if CONTINUE_KEY_4.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
                # keep track of start time/frame for later
                CONTINUE_KEY_4.frameNStart = frameN  # exact frame index
                CONTINUE_KEY_4.tStart = t  # local t and not account for scr refresh
                CONTINUE_KEY_4.tStartRefresh = tThisFlipGlobal  # on global time
                win.timeOnFlip(CONTINUE_KEY_4, 'tStartRefresh')  # time at next scr refresh
                # add timestamp to datafile
                thisExp.timestampOnFlip(win, 'CONTINUE_KEY_4.started')
                # update status
                CONTINUE_KEY_4.status = STARTED
                # keyboard checking is just starting
                waitOnFlip = True
                win.callOnFlip(CONTINUE_KEY_4.clock.reset)  # t=0 on next screen flip
            if CONTINUE_KEY_4.status == STARTED and not waitOnFlip:
                theseKeys = CONTINUE_KEY_4.getKeys(keyList=['space'], ignoreKeys=["escape"], waitRelease=False)
                _CONTINUE_KEY_4_allKeys.extend(theseKeys)
                if len(_CONTINUE_KEY_4_allKeys):
                    CONTINUE_KEY_4.keys = _CONTINUE_KEY_4_allKeys[-1].name  # just the last key pressed
                    CONTINUE_KEY_4.rt = _CONTINUE_KEY_4_allKeys[-1].rt
                    CONTINUE_KEY_4.duration = _CONTINUE_KEY_4_allKeys[-1].duration
                    # a response ends the routine
                    continueRoutine = False
            
            # *STIMULATE_IMAGE* updates
            
            # if STIMULATE_IMAGE is starting this frame...
            if STIMULATE_IMAGE.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
                # keep track of start time/frame for later
                STIMULATE_IMAGE.frameNStart = frameN  # exact frame index
                STIMULATE_IMAGE.tStart = t  # local t and not account for scr refresh
                STIMULATE_IMAGE.tStartRefresh = tThisFlipGlobal  # on global time
                win.timeOnFlip(STIMULATE_IMAGE, 'tStartRefresh')  # time at next scr refresh
                # update status
                STIMULATE_IMAGE.status = STARTED
                STIMULATE_IMAGE.setAutoDraw(True)
            
            # if STIMULATE_IMAGE is active this frame...
            if STIMULATE_IMAGE.status == STARTED:
                # update params
                pass
            
            # check for quit (typically the Esc key)
            if defaultKeyboard.getKeys(keyList=["escape"]):
                thisExp.status = FINISHED
            if thisExp.status == FINISHED or endExpNow:
                endExperiment(thisExp, win=win)
                return
            # pause experiment here if requested
            if thisExp.status == PAUSED:
                pauseExperiment(
                    thisExp=thisExp, 
                    win=win, 
                    timers=[routineTimer, globalClock], 
                    currentRoutine=STIMULATE,
                )
                # skip the frame we paused on
                continue
            
            # has a Component requested the Routine to end?
            if not continueRoutine:
                STIMULATE.forceEnded = routineForceEnded = True
            # has the Routine been forcibly ended?
            if STIMULATE.forceEnded or routineForceEnded:
                break
            # has every Component finished?
            continueRoutine = False
            for thisComponent in STIMULATE.components:
                if hasattr(thisComponent, "status") and thisComponent.status != FINISHED:
                    continueRoutine = True
                    break  # at least one component has not yet finished
            
            # refresh the screen
            if continueRoutine:  # don't flip if this routine is over or we'll get a blank screen
                win.flip()
        
        # --- Ending Routine "STIMULATE" ---
        for thisComponent in STIMULATE.components:
            if hasattr(thisComponent, "setAutoDraw"):
                thisComponent.setAutoDraw(False)
        # store stop times for STIMULATE
        STIMULATE.tStop = globalClock.getTime(format='float')
        STIMULATE.tStopRefresh = tThisFlipGlobal
        thisExp.addData('STIMULATE.stopped', STIMULATE.tStop)
        # the Routine "STIMULATE" was not non-slip safe, so reset the non-slip timer
        routineTimer.reset()
        
        # --- Prepare to start Routine "RECORD" ---
        # create an object to store info about Routine RECORD
        RECORD = data.Routine(
            name='RECORD',
            components=[SETTING_IMAGE_5, TITLE_TEXT, INPUT, TEXT, SUBMMIT_KEY],
        )
        RECORD.status = NOT_STARTED
        continueRoutine = True
        # update component parameters for each repeat
        INPUT.reset()
        # create starting attributes for SUBMMIT_KEY
        SUBMMIT_KEY.keys = []
        SUBMMIT_KEY.rt = []
        _SUBMMIT_KEY_allKeys = []
        # store start times for RECORD
        RECORD.tStartRefresh = win.getFutureFlipTime(clock=globalClock)
        RECORD.tStart = globalClock.getTime(format='float')
        RECORD.status = STARTED
        thisExp.addData('RECORD.started', RECORD.tStart)
        RECORD.maxDuration = None
        # keep track of which components have finished
        RECORDComponents = RECORD.components
        for thisComponent in RECORD.components:
            thisComponent.tStart = None
            thisComponent.tStop = None
            thisComponent.tStartRefresh = None
            thisComponent.tStopRefresh = None
            if hasattr(thisComponent, 'status'):
                thisComponent.status = NOT_STARTED
        # reset timers
        t = 0
        _timeToFirstFrame = win.getFutureFlipTime(clock="now")
        frameN = -1
        
        # --- Run Routine "RECORD" ---
        thisExp.currentRoutine = RECORD
        RECORD.forceEnded = routineForceEnded = not continueRoutine
        while continueRoutine:
            # if trial has changed, end Routine now
            if hasattr(thisTrial, 'status') and thisTrial.status == STOPPING:
                continueRoutine = False
            # get current time
            t = routineTimer.getTime()
            tThisFlip = win.getFutureFlipTime(clock=routineTimer)
            tThisFlipGlobal = win.getFutureFlipTime(clock=None)
            frameN = frameN + 1  # number of completed frames (so 0 is the first frame)
            # update/draw components on each frame
            
            # *SETTING_IMAGE_5* updates
            
            # if SETTING_IMAGE_5 is starting this frame...
            if SETTING_IMAGE_5.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
                # keep track of start time/frame for later
                SETTING_IMAGE_5.frameNStart = frameN  # exact frame index
                SETTING_IMAGE_5.tStart = t  # local t and not account for scr refresh
                SETTING_IMAGE_5.tStartRefresh = tThisFlipGlobal  # on global time
                win.timeOnFlip(SETTING_IMAGE_5, 'tStartRefresh')  # time at next scr refresh
                # update status
                SETTING_IMAGE_5.status = STARTED
                SETTING_IMAGE_5.setAutoDraw(True)
            
            # if SETTING_IMAGE_5 is active this frame...
            if SETTING_IMAGE_5.status == STARTED:
                # update params
                pass
            
            # *TITLE_TEXT* updates
            
            # if TITLE_TEXT is starting this frame...
            if TITLE_TEXT.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
                # keep track of start time/frame for later
                TITLE_TEXT.frameNStart = frameN  # exact frame index
                TITLE_TEXT.tStart = t  # local t and not account for scr refresh
                TITLE_TEXT.tStartRefresh = tThisFlipGlobal  # on global time
                win.timeOnFlip(TITLE_TEXT, 'tStartRefresh')  # time at next scr refresh
                # update status
                TITLE_TEXT.status = STARTED
                TITLE_TEXT.setAutoDraw(True)
            
            # if TITLE_TEXT is active this frame...
            if TITLE_TEXT.status == STARTED:
                # update params
                pass
            
            # *INPUT* updates
            
            # if INPUT is starting this frame...
            if INPUT.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
                # keep track of start time/frame for later
                INPUT.frameNStart = frameN  # exact frame index
                INPUT.tStart = t  # local t and not account for scr refresh
                INPUT.tStartRefresh = tThisFlipGlobal  # on global time
                win.timeOnFlip(INPUT, 'tStartRefresh')  # time at next scr refresh
                # update status
                INPUT.status = STARTED
                INPUT.setAutoDraw(True)
            
            # if INPUT is active this frame...
            if INPUT.status == STARTED:
                # update params
                pass
            
            # *TEXT* updates
            
            # if TEXT is starting this frame...
            if TEXT.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
                # keep track of start time/frame for later
                TEXT.frameNStart = frameN  # exact frame index
                TEXT.tStart = t  # local t and not account for scr refresh
                TEXT.tStartRefresh = tThisFlipGlobal  # on global time
                win.timeOnFlip(TEXT, 'tStartRefresh')  # time at next scr refresh
                # update status
                TEXT.status = STARTED
                TEXT.setAutoDraw(True)
            
            # if TEXT is active this frame...
            if TEXT.status == STARTED:
                # update params
                pass
            
            # *SUBMMIT_KEY* updates
            waitOnFlip = False
            
            # if SUBMMIT_KEY is starting this frame...
            if SUBMMIT_KEY.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
                # keep track of start time/frame for later
                SUBMMIT_KEY.frameNStart = frameN  # exact frame index
                SUBMMIT_KEY.tStart = t  # local t and not account for scr refresh
                SUBMMIT_KEY.tStartRefresh = tThisFlipGlobal  # on global time
                win.timeOnFlip(SUBMMIT_KEY, 'tStartRefresh')  # time at next scr refresh
                # add timestamp to datafile
                thisExp.timestampOnFlip(win, 'SUBMMIT_KEY.started')
                # update status
                SUBMMIT_KEY.status = STARTED
                # keyboard checking is just starting
                waitOnFlip = True
                win.callOnFlip(SUBMMIT_KEY.clock.reset)  # t=0 on next screen flip
                win.callOnFlip(SUBMMIT_KEY.clearEvents, eventType='keyboard')  # clear events on next screen flip
            if SUBMMIT_KEY.status == STARTED and not waitOnFlip:
                theseKeys = SUBMMIT_KEY.getKeys(keyList=['return'], ignoreKeys=["escape"], waitRelease=False)
                _SUBMMIT_KEY_allKeys.extend(theseKeys)
                if len(_SUBMMIT_KEY_allKeys):
                    SUBMMIT_KEY.keys = _SUBMMIT_KEY_allKeys[-1].name  # just the last key pressed
                    SUBMMIT_KEY.rt = _SUBMMIT_KEY_allKeys[-1].rt
                    SUBMMIT_KEY.duration = _SUBMMIT_KEY_allKeys[-1].duration
                    # a response ends the routine
                    continueRoutine = False
            
            # check for quit (typically the Esc key)
            if defaultKeyboard.getKeys(keyList=["escape"]):
                thisExp.status = FINISHED
            if thisExp.status == FINISHED or endExpNow:
                endExperiment(thisExp, win=win)
                return
            # pause experiment here if requested
            if thisExp.status == PAUSED:
                pauseExperiment(
                    thisExp=thisExp, 
                    win=win, 
                    timers=[routineTimer, globalClock], 
                    currentRoutine=RECORD,
                )
                # skip the frame we paused on
                continue
            
            # has a Component requested the Routine to end?
            if not continueRoutine:
                RECORD.forceEnded = routineForceEnded = True
            # has the Routine been forcibly ended?
            if RECORD.forceEnded or routineForceEnded:
                break
            # has every Component finished?
            continueRoutine = False
            for thisComponent in RECORD.components:
                if hasattr(thisComponent, "status") and thisComponent.status != FINISHED:
                    continueRoutine = True
                    break  # at least one component has not yet finished
            
            # refresh the screen
            if continueRoutine:  # don't flip if this routine is over or we'll get a blank screen
                win.flip()
        
        # --- Ending Routine "RECORD" ---
        for thisComponent in RECORD.components:
            if hasattr(thisComponent, "setAutoDraw"):
                thisComponent.setAutoDraw(False)
        # store stop times for RECORD
        RECORD.tStop = globalClock.getTime(format='float')
        RECORD.tStopRefresh = tThisFlipGlobal
        thisExp.addData('RECORD.stopped', RECORD.tStop)
        trials.addData('INPUT.text',INPUT.text)
        # check responses
        if SUBMMIT_KEY.keys in ['', [], None]:  # No response was made
            SUBMMIT_KEY.keys = None
        trials.addData('SUBMMIT_KEY.keys',SUBMMIT_KEY.keys)
        if SUBMMIT_KEY.keys != None:  # we had a response
            trials.addData('SUBMMIT_KEY.rt', SUBMMIT_KEY.rt)
            trials.addData('SUBMMIT_KEY.duration', SUBMMIT_KEY.duration)
        # the Routine "RECORD" was not non-slip safe, so reset the non-slip timer
        routineTimer.reset()
        # mark thisTrial as finished
        if hasattr(thisTrial, 'status'):
            thisTrial.status = FINISHED
        # if awaiting a pause, pause now
        if trials.status == PAUSED:
            thisExp.status = PAUSED
            pauseExperiment(
                thisExp=thisExp, 
                win=win, 
                timers=[globalClock], 
            )
            # once done pausing, restore running status
            trials.status = STARTED
        thisExp.nextEntry()
        
    # completed 1 repeats of 'trials'
    trials.status = FINISHED
    
    if thisSession is not None:
        # if running in a Session with a Liaison client, send data up to now
        thisSession.sendExperimentData()
    
    # --- Prepare to start Routine "THANKS" ---
    # create an object to store info about Routine THANKS
    THANKS = data.Routine(
        name='THANKS',
        components=[THANKS_TEXT, THANKS_DETAILED_TEXT, CONTINUE_KEY_5],
    )
    THANKS.status = NOT_STARTED
    continueRoutine = True
    # update component parameters for each repeat
    # create starting attributes for CONTINUE_KEY_5
    CONTINUE_KEY_5.keys = []
    CONTINUE_KEY_5.rt = []
    _CONTINUE_KEY_5_allKeys = []
    # store start times for THANKS
    THANKS.tStartRefresh = win.getFutureFlipTime(clock=globalClock)
    THANKS.tStart = globalClock.getTime(format='float')
    THANKS.status = STARTED
    thisExp.addData('THANKS.started', THANKS.tStart)
    THANKS.maxDuration = None
    # keep track of which components have finished
    THANKSComponents = THANKS.components
    for thisComponent in THANKS.components:
        thisComponent.tStart = None
        thisComponent.tStop = None
        thisComponent.tStartRefresh = None
        thisComponent.tStopRefresh = None
        if hasattr(thisComponent, 'status'):
            thisComponent.status = NOT_STARTED
    # reset timers
    t = 0
    _timeToFirstFrame = win.getFutureFlipTime(clock="now")
    frameN = -1
    
    # --- Run Routine "THANKS" ---
    thisExp.currentRoutine = THANKS
    THANKS.forceEnded = routineForceEnded = not continueRoutine
    while continueRoutine:
        # get current time
        t = routineTimer.getTime()
        tThisFlip = win.getFutureFlipTime(clock=routineTimer)
        tThisFlipGlobal = win.getFutureFlipTime(clock=None)
        frameN = frameN + 1  # number of completed frames (so 0 is the first frame)
        # update/draw components on each frame
        
        # *THANKS_TEXT* updates
        
        # if THANKS_TEXT is starting this frame...
        if THANKS_TEXT.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            THANKS_TEXT.frameNStart = frameN  # exact frame index
            THANKS_TEXT.tStart = t  # local t and not account for scr refresh
            THANKS_TEXT.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(THANKS_TEXT, 'tStartRefresh')  # time at next scr refresh
            # add timestamp to datafile
            thisExp.timestampOnFlip(win, 'THANKS_TEXT.started')
            # update status
            THANKS_TEXT.status = STARTED
            THANKS_TEXT.setAutoDraw(True)
        
        # if THANKS_TEXT is active this frame...
        if THANKS_TEXT.status == STARTED:
            # update params
            pass
        
        # *THANKS_DETAILED_TEXT* updates
        
        # if THANKS_DETAILED_TEXT is starting this frame...
        if THANKS_DETAILED_TEXT.status == NOT_STARTED and tThisFlip >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            THANKS_DETAILED_TEXT.frameNStart = frameN  # exact frame index
            THANKS_DETAILED_TEXT.tStart = t  # local t and not account for scr refresh
            THANKS_DETAILED_TEXT.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(THANKS_DETAILED_TEXT, 'tStartRefresh')  # time at next scr refresh
            # add timestamp to datafile
            thisExp.timestampOnFlip(win, 'THANKS_DETAILED_TEXT.started')
            # update status
            THANKS_DETAILED_TEXT.status = STARTED
            THANKS_DETAILED_TEXT.setAutoDraw(True)
        
        # if THANKS_DETAILED_TEXT is active this frame...
        if THANKS_DETAILED_TEXT.status == STARTED:
            # update params
            pass
        
        # *CONTINUE_KEY_5* updates
        
        # if CONTINUE_KEY_5 is starting this frame...
        if CONTINUE_KEY_5.status == NOT_STARTED and t >= 0.0-frameTolerance:
            # keep track of start time/frame for later
            CONTINUE_KEY_5.frameNStart = frameN  # exact frame index
            CONTINUE_KEY_5.tStart = t  # local t and not account for scr refresh
            CONTINUE_KEY_5.tStartRefresh = tThisFlipGlobal  # on global time
            win.timeOnFlip(CONTINUE_KEY_5, 'tStartRefresh')  # time at next scr refresh
            # update status
            CONTINUE_KEY_5.status = STARTED
            # keyboard checking is just starting
            CONTINUE_KEY_5.clock.reset()  # now t=0
        if CONTINUE_KEY_5.status == STARTED:
            theseKeys = CONTINUE_KEY_5.getKeys(keyList=['space'], ignoreKeys=["escape"], waitRelease=False)
            _CONTINUE_KEY_5_allKeys.extend(theseKeys)
            if len(_CONTINUE_KEY_5_allKeys):
                CONTINUE_KEY_5.keys = _CONTINUE_KEY_5_allKeys[-1].name  # just the last key pressed
                CONTINUE_KEY_5.rt = _CONTINUE_KEY_5_allKeys[-1].rt
                CONTINUE_KEY_5.duration = _CONTINUE_KEY_5_allKeys[-1].duration
                # a response ends the routine
                continueRoutine = False
        
        # check for quit (typically the Esc key)
        if defaultKeyboard.getKeys(keyList=["escape"]):
            thisExp.status = FINISHED
        if thisExp.status == FINISHED or endExpNow:
            endExperiment(thisExp, win=win)
            return
        # pause experiment here if requested
        if thisExp.status == PAUSED:
            pauseExperiment(
                thisExp=thisExp, 
                win=win, 
                timers=[routineTimer, globalClock], 
                currentRoutine=THANKS,
            )
            # skip the frame we paused on
            continue
        
        # has a Component requested the Routine to end?
        if not continueRoutine:
            THANKS.forceEnded = routineForceEnded = True
        # has the Routine been forcibly ended?
        if THANKS.forceEnded or routineForceEnded:
            break
        # has every Component finished?
        continueRoutine = False
        for thisComponent in THANKS.components:
            if hasattr(thisComponent, "status") and thisComponent.status != FINISHED:
                continueRoutine = True
                break  # at least one component has not yet finished
        
        # refresh the screen
        if continueRoutine:  # don't flip if this routine is over or we'll get a blank screen
            win.flip()
    
    # --- Ending Routine "THANKS" ---
    for thisComponent in THANKS.components:
        if hasattr(thisComponent, "setAutoDraw"):
            thisComponent.setAutoDraw(False)
    # store stop times for THANKS
    THANKS.tStop = globalClock.getTime(format='float')
    THANKS.tStopRefresh = tThisFlipGlobal
    thisExp.addData('THANKS.stopped', THANKS.tStop)
    # Run 'End Routine' code from code_4
    from openpyxl import load_workbook
    
    file_path = os.path.join(_thisDir, "conditions.xlsx")
    
    wb = load_workbook(file_path)
    ws = wb.active
    
    # 找 UP_SIDE_DOWN 所在列
    upside_col = None
    
    for cell in ws[1]:
        if cell.value == "UP_SIDE_DOWN":
            upside_col = cell.column
            break
    
    # 将整列 0/1 取反
    if upside_col is not None:
        for row in range(2, ws.max_row + 1):
            value = ws.cell(row=row, column=upside_col).value
    
            if value == 0:
                ws.cell(row=row, column=upside_col).value = 1
            elif value == 1:
                ws.cell(row=row, column=upside_col).value = 0
    
    wb.save(file_path)
    thisExp.nextEntry()
    # the Routine "THANKS" was not non-slip safe, so reset the non-slip timer
    routineTimer.reset()
    
    # mark experiment as finished
    endExperiment(thisExp, win=win)


def saveData(thisExp):
    """
    Save data from this experiment
    
    Parameters
    ==========
    thisExp : psychopy.data.ExperimentHandler
        Handler object for this experiment, contains the data to save and information about 
        where to save it to.
    """
    filename = thisExp.dataFileName
    # these shouldn't be strictly necessary (should auto-save)
    thisExp.saveAsWideText(filename + '.csv', delim='auto')
    thisExp.saveAsPickle(filename)


def endExperiment(thisExp, win=None):
    """
    End this experiment, performing final shut down operations.
    
    This function does NOT close the window or end the Python process - use `quit` for this.
    
    Parameters
    ==========
    thisExp : psychopy.data.ExperimentHandler
        Handler object for this experiment, contains the data to save and information about 
        where to save it to.
    win : psychopy.visual.Window
        Window for this experiment.
    """
    # stop any playback components
    if thisExp.currentRoutine is not None:
        for comp in thisExp.currentRoutine.getPlaybackComponents():
            comp.stop()
    if win is not None:
        # remove autodraw from all current components
        win.clearAutoDraw()
        # Flip one final time so any remaining win.callOnFlip() 
        # and win.timeOnFlip() tasks get executed
        win.flip()
    # return console logger level to WARNING
    logging.console.setLevel(logging.WARNING)
    # mark experiment handler as finished
    thisExp.status = FINISHED
    # run any 'at exit' functions
    for fcn in runAtExit:
        fcn()
    logging.flush()


def quit(thisExp, win=None, thisSession=None):
    """
    Fully quit, closing the window and ending the Python process.
    
    Parameters
    ==========
    win : psychopy.visual.Window
        Window to close.
    thisSession : psychopy.session.Session or None
        Handle of the Session object this experiment is being run from, if any.
    """
    thisExp.abort()  # or data files will save again on exit
    # make sure everything is closed down
    if win is not None:
        # Flip one final time so any remaining win.callOnFlip() 
        # and win.timeOnFlip() tasks get executed before quitting
        win.flip()
        win.close()
    logging.flush()
    if thisSession is not None:
        thisSession.stop()
    # terminate Python process
    core.quit()


# if running this experiment as a script...
if __name__ == '__main__':
    # call all functions in order
    expInfo = showExpInfoDlg(expInfo=expInfo)
    thisExp = setupData(expInfo=expInfo)
    logFile = setupLogging(filename=thisExp.dataFileName)
    win = setupWindow(expInfo=expInfo)
    setupDevices(expInfo=expInfo, thisExp=thisExp, win=win)
    run(
        expInfo=expInfo, 
        thisExp=thisExp, 
        win=win,
        globalClock='float'
    )
    saveData(thisExp=thisExp)
    quit(thisExp=thisExp, win=win)
